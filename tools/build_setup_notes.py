from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/laptcv/Desktop/LIVE")
OUTPUT = ROOT / "docs" / "LIVIN_Supabase_and_Google_Maps_Setup_Notes.docx"


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_paragraph(paragraph, fill="F1F3F4"):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def add_step(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.0
    shade_paragraph(p)
    run = p.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(9)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.right_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15

for name, size, before, after, color in (
    ("Heading 1", 20, 20, 6, "000000"),
    ("Heading 2", 16, 18, 6, "000000"),
    ("Heading 3", 14, 16, 4, "434343"),
):
    style = styles[name]
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = False
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

for name in ("List Bullet", "List Number"):
    style = styles[name]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    style.paragraph_format.left_indent = Inches(0.5)
    style.paragraph_format.first_line_indent = Inches(-0.25)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.15

title = doc.add_paragraph()
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.space_after = Pt(3)
title_run = title.add_run("L!V!N setup notes")
title_run.font.name = "Arial"
title_run.font.size = Pt(26)
title_run.font.bold = False
title_run.font.color.rgb = RGBColor(0, 0, 0)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(14)
subtitle_run = subtitle.add_run("Supabase upgrade, event photos, Venture Score, friend-only DMs, and Google Maps")
subtitle_run.font.name = "Arial"
subtitle_run.font.size = Pt(12)
subtitle_run.font.color.rgb = RGBColor(85, 85, 85)

add_body(doc, "Use these steps in order. The app code is already updated; Supabase and the Google Maps key are the two pieces you still need to configure yourself.")

doc.add_heading("What changed in the app", level=1)
for item in (
    "Events still use a vertical scroll, but each card now puts the event photo at the top and details below it.",
    "Event creation and editing can upload a cover photo. Large images are resized before upload to reduce memory crashes.",
    "Event cards prefer a real event snap as the proof image and show the number of attached real snaps.",
    "Apple map views were replaced with Google Maps SDK for iOS 10.15.0. Route buttons open Google Maps for live directions.",
    "Venture Score is stored in Supabase and awarded automatically.",
    "DMs are available only when both people follow each other. Removing either follow hides and blocks the DM.",
    "The activity calendar shows one month at a time with left and right arrows.",
):
    add_bullet(doc, item)

doc.add_heading("Part 1 — Apply the Supabase upgrade", level=1)
add_step(doc, "Open the Supabase dashboard and select the exact project used by the iOS app.")
add_step(doc, "Open SQL Editor, choose New query, and leave that browser tab open.")
add_step(doc, "On your Mac, open the file below. Select all of it and copy it.")
add_code(doc, "/Users/laptcv/Desktop/LIVE/supabase/feature_upgrade.sql")
add_step(doc, "Paste the entire file into the Supabase SQL Editor. Do not run only part of it.")
add_step(doc, "Click Run. A successful run finishes with three result sets: storage buckets, Venture Score award types, and installed RLS policies.")
add_step(doc, "If the editor reports that a policy already exists, confirm you copied the newest whole file, then run the whole file again. It is designed to be rerunnable.")

doc.add_heading("What that SQL changes", level=2)
for item in (
    "Adds profiles.venture_score and a private, idempotent Venture Score ledger.",
    "Awards points with database triggers, so users cannot award points from the iPhone client.",
    "Creates or updates event-covers, snap-media, and avatars buckets with a 5 MB JPEG limit.",
    "Uses owner_id and user-ID folders for storage ownership checks.",
    "Replaces the old DM policies with mutual-follow checks in the database.",
    "Adds explicit Data API grants needed by newer Supabase projects, while RLS still controls which rows are allowed.",
):
    add_bullet(doc, item)

doc.add_heading("Venture Score rules", level=2)
table = doc.add_table(rows=1, cols=3)
table.autofit = False
table.columns[0].width = Inches(2.55)
table.columns[1].width = Inches(1.05)
table.columns[2].width = Inches(2.9)
headers = ("Action", "Points", "Proof used")
for index, value in enumerate(headers):
    cell = table.rows[0].cells[index]
    cell.text = value
    set_cell_margins(cell)
    for run in cell.paragraphs[0].runs:
        run.bold = True
set_repeat_table_header(table.rows[0])
for action, points, proof in (
    ("Create an event", "+25", "New events row"),
    ("Join an event", "+10", "New non-owner event member"),
    ("Post a snap", "+15", "New snaps row"),
    ("Attend an event", "+25", "Snap attached to that event"),
):
    cells = table.add_row().cells
    for index, value in enumerate((action, points, proof)):
        cells[index].text = value
        set_cell_margins(cells[index])

doc.add_heading("Part 2 — Verify Supabase", level=1)
add_body(doc, "Run each check as a separate SQL Editor query.")
add_body(doc, "Check the three storage buckets:", bold_prefix="Check the three storage buckets:")
add_code(doc, "select id, public, file_size_limit, allowed_mime_types\nfrom storage.buckets\nwhere id in ('event-covers', 'snap-media', 'avatars')\norder by id;")
add_body(doc, "Expected: event-covers and avatars are public; snap-media is private; each limit is 5242880 bytes.")

add_body(doc, "Check Venture Score totals:", bold_prefix="Check Venture Score totals:")
add_code(doc, "select username, venture_score\nfrom public.profiles\norder by venture_score desc, username;")

add_body(doc, "Check the friend rule for two known usernames:", bold_prefix="Check the friend rule for two known usernames:")
add_code(doc, "select a.username as person_a, b.username as person_b,\n       private.are_friends(a.id, b.id) as are_friends\nfrom public.profiles a\ncross join public.profiles b\nwhere a.username = 'FIRST_USERNAME'\n  and b.username = 'SECOND_USERNAME';")
add_body(doc, "Expected: false until both people follow each other; true after both follows are approved.")

add_body(doc, "Finally, open Database → Advisors in Supabase and run the Security checks. Do not disable RLS to silence an error.")

doc.add_heading("Part 3 — Test event photos", level=1)
add_step(doc, "Build and sign in to the app with a real Supabase account.")
add_step(doc, "Create an event, tap the picture area, choose a photo, complete the title/location/time, and save.")
add_step(doc, "In Supabase, open Storage → event-covers. Confirm the object path begins with that user's UUID and ends in .jpg.")
add_step(doc, "Return to Events. The card should show the photo at the top. Events without a photo show the honest placeholder instead of stock imagery.")
add_step(doc, "Join the event with a second account and post a snap attached to the event. The event card should show the real-snap count after refresh.")

doc.add_heading("Part 4 — Test friend-only DMs", level=1)
add_step(doc, "Use two real accounts, A and B. Have A follow B only. Neither account should be able to start a DM yet.")
add_step(doc, "Have B follow A. The accounts are now friends, and each should appear in the other's DM list.")
add_step(doc, "Send a message in both directions. Reopen the chat to confirm history loads and Realtime updates arrive.")
add_step(doc, "Remove either follow. The old conversation should become unreadable and new message inserts should be rejected by RLS.")

doc.add_heading("Part 5 — Add the Google Maps key", level=1)
add_step(doc, "Open Google Cloud Console and select or create the project you want billed for Maps usage.")
add_step(doc, "Enable billing and enable Maps SDK for iOS.")
add_step(doc, "Create an API key. Under Application restrictions choose iOS apps, then add bundle identifier vzn.LIVE.")
add_step(doc, "Under API restrictions choose Restrict key and allow only Maps SDK for iOS. Save the key.")
add_step(doc, "In Xcode, select the LIVE project → LIVE target → Build Settings. Search for GOOGLE_MAPS_API_KEY.")
add_step(doc, "Paste the key for both Debug and Release. Do not use a server key or a Supabase key here.")
add_step(doc, "Build again. Open the Map tab, move and zoom the map, tap event/snap markers, then open an event route. The route button should open Google Maps or its web fallback.")

doc.add_heading("Fast troubleshooting", level=1)
for item in (
    "Photo upload says unauthorized: rerun feature_upgrade.sql and confirm the object path begins with the signed-in user's UUID.",
    "Photo is too large: choose another image. The app already downsizes to 2048 px and stays below the 5 MB bucket limit.",
    "Events fail after the SQL change: check Project Settings → Data API and confirm public is exposed; the upgrade script includes authenticated grants.",
    "Map says key needed: GOOGLE_MAPS_API_KEY is blank in the active build configuration.",
    "Map is blank or says the project is not authorized: enable Maps SDK for iOS, billing, and the vzn.LIVE iOS restriction on the same key.",
    "DM profile does not appear: both approved follow rows must exist, one in each direction.",
    "DM worked before but stopped after unfollow: that is intentional; the database blocks access as soon as mutual follow is gone.",
):
    add_bullet(doc, item)

doc.add_heading("Files that matter", level=1)
for path, purpose in (
    ("supabase/feature_upgrade.sql", "The one SQL file to run in the dashboard."),
    ("LIVE/GoogleMapsSupport.swift", "Google Maps configuration and SwiftUI bridge."),
    ("LIVE/EventService.swift", "Event cover upload and real-snap proof image loading."),
    ("LIVE/DMService.swift", "Friend-only DM profile loading and conversation creation."),
    ("LIVE/UploadImageEncoder.swift", "Upload resizing and compression."),
):
    add_bullet(doc, f"{path} — {purpose}")

doc.add_heading("Official references", level=1)
for reference in (
    "Supabase RLS: https://supabase.com/docs/guides/database/postgres/row-level-security",
    "Supabase Storage access control: https://supabase.com/docs/guides/storage/security/access-control",
    "Google Maps iOS setup: https://developers.google.com/maps/documentation/ios-sdk/config",
    "Google Maps SwiftUI integration: https://developers.google.com/maps/documentation/ios-sdk/map",
):
    add_bullet(doc, reference)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
